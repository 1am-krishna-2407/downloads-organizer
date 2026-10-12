import hashlib
import json
import os
from contextlib import closing
from dataclasses import replace
from pathlib import Path

import pytest

from downloads_organizer import database, decision_engine, organizer, pipeline, reporter
from downloads_organizer.config import ExtractionSettings
from downloads_organizer.decision_engine import ARCHIVE, CLEANUP, KEEP, LEAVE, REVIEW
from downloads_organizer.scanner import scan

AI = {
    "document_type": "Reference material", "category": "Reference", "subject": "x", "summary": "A book.",
    "purpose": "x", "importance": "medium", "recommended_action": "ARCHIVE", "confidence": 0.9,
    "reasons": ["It is a book"], "sensitivity": "none", "duplicate_suspected": False,
}


@pytest.fixture
def env(config, downloads):
    """Downloads with documents, duplicates and a code project; processed but not yet classified."""
    for name, text in {
        "payslip_march.txt": "Payslip March. Net pay 50000",
        "book.txt": "Chapter 1. Algorithms",
        "ticket.txt": "Train ticket for a past journey",
        "unsure.txt": "Some text",
        "scam.txt": "Bank statement",
        "notes.txt": "Lecture notes week 1",
        "notes (1).txt": "Lecture notes week 1",
    }.items():
        (downloads / name).write_text(text, encoding="utf-8")
    project = downloads / "Tetris-master"
    project.mkdir()
    (project / "README.txt").write_text("Tetris readme", encoding="utf-8")
    other_project = downloads / "myapp"
    (other_project / "docs").mkdir(parents=True)
    (other_project / "package.json").write_text("{}", encoding="utf-8")
    (other_project / "docs" / "guide.txt").write_text("App guide", encoding="utf-8")

    cfg = replace(config, extraction=ExtractionSettings(ocr_enabled=False))
    conn = database.connect(cfg.database_path)
    database.save_scan(conn, scan(cfg))
    pipeline.process_documents(conn, cfg)
    yield cfg, conn, downloads
    conn.close()


def classify(conn, filename, **changes):
    payload = {"classification": {**AI, **changes}, "model_used": "m", "prompt_version": "v"}
    conn.execute(
        "UPDATE documents SET classification_status = 'done', classification_json = ?, document_type = ?,"
        " recommended_action = ?, confidence = ? WHERE filename = ?",
        (json.dumps(payload), payload["classification"]["document_type"],
         payload["classification"]["recommended_action"], payload["classification"]["confidence"], filename),
    )
    conn.commit()


def decisions_by_name(conn, cfg):
    decision_engine.decide_all(conn, cfg)
    return {r["filename"]: r for r in conn.execute(
        "SELECT d.filename, x.* FROM decisions x JOIN documents d ON d.id = x.document_id")}


def test_rules(env):
    cfg, conn, _ = env
    classify(conn, "payslip_march.txt", document_type="Financial document", recommended_action="ARCHIVE",
             sensitivity="financial")
    classify(conn, "book.txt")
    classify(conn, "ticket.txt", document_type="Temporary or miscellaneous download",
             recommended_action="CLEANUP_CANDIDATE", confidence=0.9)
    classify(conn, "unsure.txt", confidence=0.4)
    classify(conn, "scam.txt", document_type="Reference material", recommended_action="CLEANUP_CANDIDATE",
             sensitivity="financial", confidence=0.95)
    classify(conn, "notes.txt", document_type="Academic notes", recommended_action="KEEP")
    classify(conn, "notes (1).txt", document_type="Academic notes", recommended_action="KEEP")
    d = decisions_by_name(conn, cfg)

    # Protected types are kept whatever the AI says, and filed by type.
    assert (d["payslip_march.txt"]["action"], d["payslip_march.txt"]["destination"]) == (KEEP, "Finance")
    assert "PROTECTED_TYPE" in d["payslip_march.txt"]["rules"]
    assert (d["book.txt"]["action"], d["book.txt"]["destination"]) == (ARCHIVE, "Archive/Reference")
    assert d["ticket.txt"]["action"] == CLEANUP
    assert d["unsure.txt"]["action"] == REVIEW and "LOW_CONFIDENCE" in d["unsure.txt"]["rules"]
    assert d["scam.txt"]["action"] == REVIEW and "SENSITIVE_NOT_CLEANUP" in d["scam.txt"]["rules"]

    # One copy of a duplicate is kept; the "(1)" copy becomes a cleanup candidate.
    assert (d["notes.txt"]["action"], d["notes.txt"]["destination"]) == (KEEP, "Education/Notes")
    assert d["notes (1).txt"]["action"] == CLEANUP
    assert d["notes (1).txt"]["destination"] == decision_engine.DUPLICATES_DIR

    # Project folders (by name or by marker file) aren't even scanned.
    assert "README.txt" not in d and "guide.txt" not in d
    # Non-documents and unclassified documents stay where they are.
    assert d["setup.exe"]["action"] == LEAVE
    assert d["invoice.txt"]["action"] == LEAVE and "NOT_CLASSIFIED" in d["invoice.txt"]["rules"]
    # Unreadable documents go to review.
    assert d["Notes.DOCX"]["action"] == REVIEW and "UNREADABLE" in d["Notes.DOCX"]["rules"]
    # Every decision explains itself.
    assert all(json.loads(r["reasons"]) for r in d.values())


def test_project_files_are_left_in_place_even_if_scanned(env):
    cfg, conn, _ = env
    cfg = replace(cfg, skip_project_folders=False)
    database.save_scan(conn, scan(cfg))
    pipeline.process_documents(conn, cfg)
    d = decisions_by_name(conn, cfg)
    assert d["README.txt"]["action"] == LEAVE and "IN_PROJECT_FOLDER" in d["README.txt"]["rules"]
    assert d["guide.txt"]["action"] == LEAVE


def test_photos_are_left_alone_but_text_images_are_reviewed(env, monkeypatch):
    from downloads_organizer import extractors

    cfg, conn, downloads = env
    cfg = replace(cfg, supported_extensions=cfg.supported_extensions | {".png"})
    (downloads / "holiday.png").write_bytes(b"\x89PNG photo")
    (downloads / "scan.png").write_bytes(b"\x89PNG scan")
    monkeypatch.setattr(extractors, "ocr_available", lambda settings: True)
    monkeypatch.setitem(extractors.EXTRACTORS, ".png", lambda path, settings: (
        extractors.Extraction(extractors.EMPTY, method="ocr") if path.name == "holiday.png"
        else extractors.Extraction(extractors.ERROR, error="broken")))
    database.save_scan(conn, scan(cfg))
    pipeline.process_documents(conn, cfg)
    d = decisions_by_name(conn, cfg)
    assert d["holiday.png"]["action"] == LEAVE and "PHOTO" in d["holiday.png"]["rules"]
    assert d["scan.png"]["action"] == REVIEW


def test_similar_versions_are_noted(env):
    cfg, conn, downloads = env
    body = " ".join(f"Experience item {i}: built data pipelines and APIs." for i in range(60))
    (downloads / "cv_v1.txt").write_text("Jane Doe CV. " + body, encoding="utf-8")
    (downloads / "cv_v2.txt").write_text("Jane Doe CV updated. " + body + " New role added.", encoding="utf-8")
    database.save_scan(conn, scan(cfg))
    stats = pipeline.process_documents(conn, cfg)
    assert stats.similar_groups == 1
    classify(conn, "cv_v1.txt", document_type="Resume or CV")
    classify(conn, "cv_v2.txt", document_type="Resume or CV")
    d = decisions_by_name(conn, cfg)
    assert "SIMILAR_VERSIONS" in d["cv_v1.txt"]["rules"]
    assert "cv_v2.txt" in d["cv_v1.txt"]["reasons"]


def test_user_override_wins(env):
    cfg, conn, _ = env
    classify(conn, "book.txt")
    doc_id = conn.execute("SELECT id FROM documents WHERE filename = 'book.txt'").fetchone()[0]
    conn.execute("INSERT INTO user_overrides (document_id, action, destination, updated_at)"
                 " VALUES (?, 'KEEP', 'Education/Notes', 'now')", (doc_id,))
    d = decisions_by_name(conn, cfg)
    assert (d["book.txt"]["action"], d["book.txt"]["destination"]) == (KEEP, "Education/Notes")


def test_review_report(env):
    from openpyxl import load_workbook

    cfg, conn, _ = env
    classify(conn, "payslip_march.txt", document_type="Financial document", recommended_action="ARCHIVE")
    decision_engine.decide_all(conn, cfg)
    wb = load_workbook(reporter.write_review_report(conn, cfg.reports_dir))
    assert wb.sheetnames == ["Summary", "Proposed Moves", "Review Queue", "Cleanup Candidates",
                             "Rule Overrides", "Left In Place"]
    overrides = list(wb["Rule Overrides"].iter_rows(min_row=2, values_only=True))
    assert overrides[0][:3] == ("payslip_march.txt", "ARCHIVE", "KEEP")


def _snapshot(folder: Path) -> dict[str, str]:
    return {str(p.relative_to(folder)): hashlib.sha256(p.read_bytes()).hexdigest()
            for p in folder.rglob("*") if p.is_file()}


def test_plan_is_read_only_and_apply_then_rollback_restores_everything(env):
    cfg, conn, downloads = env
    classify(conn, "payslip_march.txt", document_type="Financial document")
    classify(conn, "book.txt")
    classify(conn, "ticket.txt", document_type="Temporary or miscellaneous download",
             recommended_action="CLEANUP_CANDIDATE")
    decision_engine.decide_all(conn, cfg)
    before = _snapshot(downloads)

    moves, skipped = organizer.plan_moves(conn, cfg)
    assert _snapshot(downloads) == before          # planning moves nothing
    assert {m.source.name for m in moves} == {"payslip_march.txt", "book.txt"}  # cleanup not included by default
    assert skipped == []

    batch, moved, failures = organizer.apply_moves(conn, moves)
    assert (moved, failures) == (2, [])
    assert (cfg.organized_dir / "Finance" / "payslip_march.txt").read_text(encoding="utf-8").startswith("Payslip")
    assert not (downloads / "book.txt").exists()
    ops = conn.execute("SELECT status FROM file_operations WHERE batch_id = ?", (batch,)).fetchall()
    assert [o[0] for o in ops] == ["completed", "completed"]

    # Moved files are not proposed again.
    decision_engine.decide_all(conn, cfg)
    assert organizer.plan_moves(conn, cfg)[0] == []

    restored, problems = organizer.rollback(conn, cfg, batch)
    assert (restored, problems) == (2, [])
    assert _snapshot(downloads) == before
    assert not (cfg.organized_dir / "Finance").exists()   # emptied folders are tidied away


def test_never_overwrites_existing_files(env):
    cfg, conn, downloads = env
    classify(conn, "book.txt")
    decision_engine.decide_all(conn, cfg)
    existing = cfg.organized_dir / "Archive" / "Reference" / "book.txt"
    existing.parent.mkdir(parents=True)
    existing.write_text("someone else's file", encoding="utf-8")

    moves, _ = organizer.plan_moves(conn, cfg)
    assert moves[0].destination.name == "book (2).txt"
    organizer.apply_moves(conn, moves)
    assert existing.read_text(encoding="utf-8") == "someone else's file"
    assert (existing.parent / "book (2).txt").read_text(encoding="utf-8") == "Chapter 1. Algorithms"


def test_two_files_with_same_name_get_distinct_destinations(env):
    cfg, conn, downloads = env
    (downloads / "nested" / "book.txt").write_text("A different book", encoding="utf-8")
    database.save_scan(conn, scan(cfg))
    pipeline.process_documents(conn, cfg)
    classify(conn, "book.txt")
    decision_engine.decide_all(conn, cfg)
    moves, _ = organizer.plan_moves(conn, cfg)
    assert sorted(m.destination.name for m in moves) == ["book (2).txt", "book.txt"]


def test_changed_or_missing_files_are_skipped(env):
    cfg, conn, downloads = env
    classify(conn, "book.txt")
    classify(conn, "payslip_march.txt", document_type="Financial document")
    decision_engine.decide_all(conn, cfg)
    (downloads / "book.txt").write_text("edited after analysis", encoding="utf-8")
    os.utime(downloads / "book.txt", (1_900_000_000, 1_900_000_000))
    (downloads / "payslip_march.txt").unlink()

    moves, skipped = organizer.plan_moves(conn, cfg)
    assert moves == []
    assert sorted(s.reason for s in skipped) == [
        "File changed since it was analysed; run again", "File no longer exists; rescan"]


def test_failed_move_is_logged_and_others_continue(env, monkeypatch):
    cfg, conn, downloads = env
    classify(conn, "book.txt")
    classify(conn, "payslip_march.txt", document_type="Financial document")
    decision_engine.decide_all(conn, cfg)
    moves, _ = organizer.plan_moves(conn, cfg)

    real = organizer._move_verified

    def flaky(source, destination):
        if source.name == "book.txt":
            raise PermissionError(13, "File is open in another program")
        real(source, destination)

    monkeypatch.setattr(organizer, "_move_verified", flaky)
    _, moved, failures = organizer.apply_moves(conn, moves)
    assert moved == 1 and len(failures) == 1
    assert (downloads / "book.txt").exists()
    statuses = dict(conn.execute("SELECT source_path, status FROM file_operations").fetchall())
    assert statuses[str(downloads / "book.txt")] == "failed"


def test_rollback_does_not_overwrite_a_new_file_at_the_original_path(env):
    cfg, conn, downloads = env
    classify(conn, "book.txt")
    decision_engine.decide_all(conn, cfg)
    batch, _, _ = organizer.apply_moves(conn, organizer.plan_moves(conn, cfg)[0])
    (downloads / "book.txt").write_text("a new download with the same name", encoding="utf-8")

    restored, problems = organizer.rollback(conn, cfg, batch)
    assert restored == 0 and len(problems) == 1
    assert (downloads / "book.txt").read_text(encoding="utf-8") == "a new download with the same name"
    assert (cfg.organized_dir / "Archive" / "Reference" / "book.txt").exists()


def test_cli_organize_preview_apply_and_undo(env, capsys):
    from downloads_organizer.cli import main

    cfg, conn, downloads = env
    classify(conn, "book.txt")
    cfg_file = cfg.data_dir.parent / "config.yaml"
    cfg_file.write_text(
        f"downloads_dir: '{downloads.as_posix()}'\nsupported_extensions: [.txt]\n"
        f"data_dir: '{cfg.data_dir.as_posix()}'\nreports_dir: '{cfg.reports_dir.as_posix()}'\n"
        "automation: {min_file_age_seconds: 0}\n", encoding="utf-8")

    assert main(["organize", "--config", str(cfg_file)]) == 0
    assert "Preview only" in capsys.readouterr().out
    assert (downloads / "book.txt").exists()

    assert main(["organize", "--config", str(cfg_file), "--apply"]) == 1  # no TTY, no --yes: refuses
    assert (downloads / "book.txt").exists()

    assert main(["organize", "--config", str(cfg_file), "--apply", "--yes"]) == 0
    assert not (downloads / "book.txt").exists()
    assert main(["organize", "--config", str(cfg_file), "--rollback", "last", "--yes"]) == 0
    assert (downloads / "book.txt").exists()
