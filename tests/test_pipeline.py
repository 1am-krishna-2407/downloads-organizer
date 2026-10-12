import os
from contextlib import closing
from dataclasses import replace

import pytest
from openpyxl import load_workbook

from downloads_organizer import database, extractors, pipeline, reporter
from downloads_organizer.cli import main
from downloads_organizer.config import AutomationSettings, ExtractionSettings
from downloads_organizer.scanner import scan
from tests.test_scanner import _snapshot


@pytest.fixture
def setup(config, downloads):
    """Downloads with real duplicates, mapped to a config with OCR switched off."""
    (downloads / "resume.pdf").unlink()
    (downloads / "resume.txt").write_text("Jane Doe resume: Python, SQL", encoding="utf-8")
    (downloads / "resume (1).txt").write_text("Jane Doe resume: Python, SQL", encoding="utf-8")
    (downloads / "nested" / "resume copy.txt").write_text("Jane Doe resume: Python, SQL", encoding="utf-8")
    (downloads / "movie.mkv").write_bytes(b"m" * 4096)       # unique size, not a document
    (downloads / "setup (1).exe").write_bytes(b"MZ")          # same bytes as setup.exe
    (downloads / "blank.txt").write_bytes(b"")
    (downloads / "blank2.txt").write_bytes(b"")
    cfg = replace(
        config,
        supported_extensions=config.supported_extensions | {".png"},
        extraction=ExtractionSettings(ocr_enabled=False),
    )
    return cfg, downloads


def run(cfg):
    conn = database.connect(cfg.database_path)
    database.save_scan(conn, scan(cfg))
    stats = pipeline.process_documents(conn, cfg)
    return conn, stats


def docs(conn):
    return {r["filename"]: r for r in conn.execute("SELECT * FROM documents")}


def test_hashes_extracts_and_groups_duplicates(setup):
    cfg, _ = setup
    conn, stats = run(cfg)
    with closing(conn):
        rows = docs(conn)
        assert rows["resume.txt"]["extraction_status"] == extractors.OK
        assert rows["setup.exe"]["extraction_status"] == pipeline.NOT_APPLICABLE

        # Identical content is extracted once and reused for the copies.
        assert stats.extracted == 5  # 7 readable files, but the 3 resume copies share one extraction
        assert conn.execute("SELECT COUNT(*) FROM extracted_text WHERE text LIKE 'Jane Doe%'").fetchone()[0] == 1

        groups = {}
        for group_id, filename in conn.execute(
            "SELECT g.group_id, d.filename FROM duplicate_groups g JOIN documents d ON d.id = g.document_id"
        ):
            groups.setdefault(group_id, set()).add(filename)
        assert sorted(groups.values(), key=len) == [
            {"setup.exe", "setup (1).exe"},
            {"resume.txt", "resume (1).txt", "resume copy.txt"},
        ]
        # Zero-byte files are not reported as duplicates.
        assert stats.duplicate_groups == 2


def test_non_documents_with_unique_size_are_not_hashed(setup):
    cfg, _ = setup
    conn, _ = run(cfg)
    with closing(conn):
        rows = docs(conn)
        assert rows["movie.mkv"]["sha256"] is None
        assert rows["setup.exe"]["sha256"] is not None


def test_second_run_does_no_work_and_changes_are_reprocessed(setup):
    cfg, downloads = setup
    conn, _ = run(cfg)
    conn.close()

    conn, stats = run(cfg)
    with closing(conn):
        assert (stats.hashed, stats.extracted, stats.extraction_cache_hits) == (0, 0, 0)

    changed = downloads / "resume (1).txt"
    changed.write_text("Jane Doe resume v2: Python, SQL, Go", encoding="utf-8")
    os.utime(changed, (1_800_000_000, 1_800_000_000))
    conn, stats = run(cfg)
    with closing(conn):
        assert (stats.hashed, stats.extracted) == (1, 1)
        assert stats.duplicate_groups == 2
        group = conn.execute(
            "SELECT COUNT(*) FROM duplicate_groups g JOIN documents d ON d.id = g.document_id"
            " WHERE d.filename LIKE 'resume%'"
        ).fetchone()[0]
        assert group == 2


def test_unreadable_files_go_to_review(setup):
    cfg, downloads = setup
    (downloads / "broken.pdf").write_bytes(b"%PDF-1.4 garbage")
    (downloads / "scan.png").write_bytes(b"\x89PNG\r\n\x1a\n")
    conn, _ = run(cfg)
    with closing(conn):
        rows = docs(conn)
        assert rows["broken.pdf"]["extraction_status"] == extractors.ERROR
        assert rows["broken.pdf"]["processing_error"]
        assert rows["scan.png"]["extraction_status"] == extractors.NEEDS_OCR

        paths = reporter.write_document_report(conn, cfg.reports_dir)
        wb = load_workbook(next(p for p in paths if p.suffix == ".xlsx"))
        assert wb.sheetnames == ["Summary", "Duplicates", "Needs Review", "Documents"]
        review = {row[1] for row in wb["Needs Review"].iter_rows(min_row=2, values_only=True)}
        assert {"broken.pdf", "scan.png", "blank.txt", "blank2.txt"} <= review
        assert "setup.exe" not in review


def test_unhashable_file_is_recorded_and_run_continues(setup, monkeypatch):
    cfg, _ = setup
    real = pipeline.sha256_file

    def flaky(path):
        if path.endswith("invoice.txt"):
            raise PermissionError(13, "Access is denied")
        return real(path)

    monkeypatch.setattr(pipeline, "sha256_file", flaky)
    conn, stats = run(cfg)
    with closing(conn):
        rows = docs(conn)
        assert rows["invoice.txt"]["extraction_status"] == extractors.ERROR
        assert "PermissionError" in rows["invoice.txt"]["processing_error"]
        assert rows["resume.txt"]["extraction_status"] == extractors.OK
        assert stats.errors >= 1


def test_needs_ocr_is_retried_once_ocr_becomes_available(setup, monkeypatch):
    cfg, downloads = setup
    (downloads / "scan.png").write_bytes(b"\x89PNG\r\n\x1a\n")
    conn, _ = run(cfg)
    conn.close()

    monkeypatch.setattr(extractors, "ocr_available", lambda settings: True)
    monkeypatch.setitem(extractors.EXTRACTORS, ".png", lambda path, settings: extractors.Extraction(
        extractors.OK, text="Scanned certificate", method="ocr"))
    conn, stats = run(cfg)
    with closing(conn):
        assert stats.extracted == 1
        assert docs(conn)["scan.png"]["extraction_status"] == extractors.OK


def test_processing_does_not_modify_downloads(setup):
    cfg, downloads = setup
    before = _snapshot(downloads)
    conn, _ = run(cfg)
    conn.close()
    assert _snapshot(downloads) == before


def test_migrates_phase1_database(tmp_path):
    import sqlite3

    path = tmp_path / "old.db"
    phase1_schema = database.SCHEMA.replace("last_scan_run         INTEGER,\n    processing_error      TEXT", "last_scan_run         INTEGER")
    assert phase1_schema != database.SCHEMA
    old = sqlite3.connect(path)
    old.executescript(phase1_schema)
    old.close()
    with closing(database.connect(path)) as conn:
        columns = {r["name"] for r in conn.execute("PRAGMA table_info(documents)")}
        assert "processing_error" in columns


def test_cli_process(tmp_path, downloads, capsys):
    cfg = tmp_path / "config.yaml"
    cfg.write_text(
        f"downloads_dir: '{downloads.as_posix()}'\n"
        "data_dir: data\nreports_dir: reports\n"
        "supported_extensions: [.pdf, .txt]\n"
        "extraction: {ocr_enabled: false}\n",
        encoding="utf-8",
    )
    assert main(["process", "--config", str(cfg)]) == 0
    out = capsys.readouterr().out
    assert "Duplicate groups:" in out
    assert list((tmp_path / "reports").glob("documents_*.xlsx"))


def test_recently_modified_files_wait_for_the_next_run(setup):
    cfg, downloads = setup
    os.utime(downloads / "invoice.txt" if (downloads / "invoice.txt").exists() else downloads / "nested" / "invoice.txt",
             (1_700_000_000, 1_700_000_000))
    cfg = replace(cfg, automation=AutomationSettings(min_file_age_seconds=3600))
    conn, stats = run(cfg)
    with closing(conn):
        rows = docs(conn)
        assert rows["invoice.txt"]["extraction_status"] == extractors.OK   # old file: processed
        assert rows["resume.txt"]["extraction_status"] == "pending"        # just written: deferred
        assert stats.deferred > 0


def test_parallel_extraction_matches_serial(setup, monkeypatch):
    cfg, downloads = setup
    for i in range(12):
        (downloads / f"note_{i}.txt").write_text(f"Lecture note number {i} about databases", encoding="utf-8")
    monkeypatch.setattr(pipeline, "PARALLEL_MIN_JOBS", 2)
    parallel = replace(cfg, extraction=replace(cfg.extraction, workers=2))
    conn, stats = run(parallel)
    with closing(conn):
        rows = docs(conn)
        assert all(rows[f"note_{i}.txt"]["extraction_status"] == extractors.OK for i in range(12))
        text = conn.execute("SELECT t.text FROM documents d JOIN extracted_text t ON t.sha256 = d.sha256"
                            " WHERE d.filename = 'note_7.txt'").fetchone()[0]
        assert text == "Lecture note number 7 about databases"


def test_crashed_worker_falls_back_to_main_process(setup, monkeypatch):
    from concurrent.futures.process import BrokenProcessPool

    cfg, downloads = setup

    class CrashingPool:
        def __init__(self, *a, **k): pass
        def __enter__(self): raise BrokenProcessPool("worker died")
        def __exit__(self, *a): return False

    monkeypatch.setattr(pipeline, "ProcessPoolExecutor", CrashingPool)
    monkeypatch.setattr(pipeline, "PARALLEL_MIN_JOBS", 1)
    conn, stats = run(replace(cfg, extraction=replace(cfg.extraction, workers=2)))
    with closing(conn):
        assert docs(conn)["resume.txt"]["extraction_status"] == extractors.OK
