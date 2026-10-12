import csv
import os
from contextlib import closing

from openpyxl import load_workbook

from downloads_organizer import database, reporter
from downloads_organizer.cli import main
from downloads_organizer.scanner import scan


def test_save_scan_tracks_changes_and_missing_files(config, downloads):
    with closing(database.connect(config.database_path)) as conn:
        database.save_scan(conn, scan(config))
        assert conn.execute("SELECT COUNT(*) FROM documents").fetchone()[0] == 6

        # Simulate later phases having processed these files.
        conn.execute(
            "UPDATE documents SET sha256 = 'abc', extraction_status = 'ok', classification_status = 'done'"
        )
        conn.commit()

        changed = downloads / "nested" / "invoice.txt"
        changed.write_text("Invoice 42 - revised", encoding="utf-8")
        os.utime(changed, (1_800_000_000, 1_800_000_000))
        (downloads / "setup.exe").unlink()
        database.save_scan(conn, scan(config))

        rows = {r["filename"]: r for r in conn.execute("SELECT * FROM documents")}
        assert rows["invoice.txt"]["sha256"] is None
        assert rows["invoice.txt"]["classification_status"] == "pending"
        assert rows["resume.pdf"]["sha256"] == "abc"
        assert rows["resume.pdf"]["classification_status"] == "done"
        assert rows["setup.exe"]["present"] == 0
        assert rows["resume.pdf"]["present"] == 1
        assert conn.execute("SELECT COUNT(*) FROM scan_runs").fetchone()[0] == 2


def test_write_inventory_csv_and_xlsx(config):
    result = scan(config)
    paths = reporter.write_inventory(result, config.reports_dir)
    csv_path = next(p for p in paths if p.suffix == ".csv")
    xlsx_path = next(p for p in paths if p.suffix == ".xlsx")

    with open(csv_path, encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh))
    assert len(rows) == 6
    assert {r["filename"] for r in rows} >= {"café menu.xlsx", "resume.pdf"}

    wb = load_workbook(xlsx_path)
    assert wb.sheetnames == ["Summary", "Inventory", "Errors"]
    assert wb["Inventory"].max_row == 7


def test_human_size():
    assert reporter.human_size(512) == "512 B"
    assert reporter.human_size(2048) == "2.0 KB"
    assert reporter.human_size(5 * 1024**3) == "5.0 GB"


def test_cli_scan(tmp_path, downloads, capsys):
    cfg = tmp_path / "config.yaml"
    cfg.write_text(
        f"downloads_dir: '{downloads.as_posix()}'\n"
        "data_dir: data\nreports_dir: reports\n"
        "supported_extensions: [.pdf]\n"
        "skip_extensions: [.crdownload]\n",
        encoding="utf-8",
    )
    assert main(["scan", "--config", str(cfg)]) == 0
    out = capsys.readouterr().out
    assert "Files inventoried: 6 (1 supported)" in out  # packaged defaults skip desktop.ini and ~$ files
    assert (tmp_path / "data" / "organizer.db").exists()
    assert list((tmp_path / "reports").glob("inventory_*.xlsx"))
