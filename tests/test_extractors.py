from pathlib import Path

import docx
import openpyxl
import pymupdf
import pytest
from pptx import Presentation

from downloads_organizer import extractors
from downloads_organizer.config import ExtractionSettings
from downloads_organizer.extractors import extract

SETTINGS = ExtractionSettings(ocr_enabled=False)


def make_pdf(path: Path, pages: list[str], **save_kwargs) -> Path:
    doc = pymupdf.open()
    for text in pages:
        page = doc.new_page()
        if text:
            page.insert_text((72, 72), text)
    doc.save(path, **save_kwargs)
    doc.close()
    return path


def test_pdf(tmp_path):
    path = make_pdf(tmp_path / "a.pdf", ["Curriculum Vitae of Jane Doe, software engineer", "Education: BSc"])
    result = extract(path, SETTINGS)
    assert result.status == extractors.OK
    assert "Curriculum Vitae" in result.text and "Education" in result.text
    assert result.page_count == 2
    assert result.method == "pymupdf"
    assert not result.truncated


def test_pdf_page_limit_marks_truncated(tmp_path):
    path = make_pdf(tmp_path / "long.pdf", [f"This is page number {i} of the report" for i in range(5)])
    result = extract(path, ExtractionSettings(ocr_enabled=False, max_pages=2))
    assert result.truncated
    assert "page number 1" in result.text and "page number 3" not in result.text


def test_pdf_without_text_needs_ocr_when_tesseract_unavailable(tmp_path):
    path = make_pdf(tmp_path / "scan.pdf", ["", ""])
    assert extract(path, SETTINGS).status == extractors.NEEDS_OCR


def test_encrypted_pdf(tmp_path):
    path = make_pdf(
        tmp_path / "locked.pdf", ["secret"],
        encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="user",
    )
    assert extract(path, SETTINGS).status == extractors.ENCRYPTED


def test_corrupt_pdf_is_an_error_not_an_exception(tmp_path):
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"%PDF-1.4 this is not really a pdf")
    result = extract(path, SETTINGS)
    assert result.status == extractors.ERROR
    assert result.error


def test_docx_paragraphs_and_tables(tmp_path):
    document = docx.Document()
    document.add_paragraph("Offer letter for the position of Data Analyst")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Salary"
    table.rows[0].cells[1].text = "100"
    path = tmp_path / "offer.docx"
    document.save(path)

    result = extract(path, SETTINGS)
    assert result.status == extractors.OK
    assert "Offer letter" in result.text and "Salary\t100" in result.text


def test_password_protected_office_file(tmp_path):
    path = tmp_path / "locked.docx"
    path.write_bytes(extractors.OLE_MAGIC + b"\x00" * 100)
    assert extract(path, SETTINGS).status == extractors.ENCRYPTED


def test_pptx(tmp_path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Quarterly project review"
    path = tmp_path / "deck.pptx"
    presentation.save(path)

    result = extract(path, SETTINGS)
    assert result.status == extractors.OK
    assert "Quarterly project review" in result.text
    assert result.page_count == 1


def test_xlsx(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Expenses"
    ws.append(["Item", "Amount"])
    ws.append(["Rent", 1200])
    path = tmp_path / "budget.xlsx"
    wb.save(path)

    result = extract(path, SETTINGS)
    assert result.status == extractors.OK
    assert "[Sheet: Expenses]" in result.text and "Rent\t1200" in result.text


@pytest.mark.parametrize(
    "data, expected",
    [
        ("Café résumé".encode("utf-8"), "Café résumé"),
        ("﻿With BOM".encode("utf-8"), "With BOM"),
        ("UTF-16 text".encode("utf-16"), "UTF-16 text"),
        ("Windows café".encode("cp1252"), "Windows café"),
    ],
)
def test_text_encodings(tmp_path, data, expected):
    path = tmp_path / "notes.txt"
    path.write_bytes(data)
    result = extract(path, SETTINGS)
    assert result.status == extractors.OK
    assert result.text == expected


def test_text_truncated_at_max_chars(tmp_path):
    path = tmp_path / "big.txt"
    path.write_text("é" * 50, encoding="utf-8")
    result = extract(path, ExtractionSettings(ocr_enabled=False, max_chars=10))
    assert result.text == "é" * 10
    assert result.truncated


def test_whitespace_is_normalised(tmp_path):
    path = tmp_path / "messy.txt"
    path.write_text("a  \t b\n\n\n\nc\x00", encoding="utf-8")
    assert extract(path, SETTINGS).text == "a\tb\n\nc"


def test_empty_text_file(tmp_path):
    path = tmp_path / "empty.txt"
    path.write_bytes(b"   \n ")
    assert extract(path, SETTINGS).status == extractors.EMPTY


def test_image_needs_ocr_when_tesseract_unavailable(tmp_path):
    path = tmp_path / "photo.png"
    path.write_bytes(b"\x89PNG\r\n\x1a\n")
    assert extract(path, SETTINGS).status == extractors.NEEDS_OCR


def test_legacy_doc_is_unsupported(tmp_path):
    path = tmp_path / "old.doc"
    path.write_bytes(b"whatever")
    assert extract(path, SETTINGS).status == extractors.UNSUPPORTED


def test_too_large(tmp_path):
    path = tmp_path / "huge.txt"
    path.write_bytes(b"x" * (1024 * 1024 + 1))
    assert extract(path, ExtractionSettings(ocr_enabled=False, max_file_mb=1)).status == extractors.TOO_LARGE


def test_extraction_does_not_modify_files(tmp_path):
    path = make_pdf(tmp_path / "a.pdf", ["Some text on a page"])
    before = (path.read_bytes(), path.stat().st_mtime_ns)
    extract(path, SETTINGS)
    assert (path.read_bytes(), path.stat().st_mtime_ns) == before
