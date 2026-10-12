"""Read-only text extraction for PDF, Word, PowerPoint, spreadsheets, text and images (Phase 2).

Every extractor returns an ``Extraction`` with a status instead of raising, so
one bad file never stops a run. Document text is never logged.
"""

from __future__ import annotations

import logging
import os
import re
import shutil
import sys
from contextlib import contextmanager
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path

from downloads_organizer.config import ExtractionSettings

log = logging.getLogger(__name__)

# Bump when extraction logic changes so cached text is regenerated.
EXTRACTOR_VERSION = "2"

OK = "ok"
EMPTY = "empty"              # readable, but no text found
NEEDS_OCR = "needs_ocr"      # scanned PDF or image, and OCR is unavailable
ENCRYPTED = "encrypted"
TOO_LARGE = "too_large"
UNSUPPORTED = "unsupported"  # type we cannot read (e.g. legacy .doc)
ERROR = "error"              # corrupt or unreadable

# Statuses that mean a person should look at the file.
REVIEW_STATUSES = frozenset({EMPTY, NEEDS_OCR, ENCRYPTED, TOO_LARGE, UNSUPPORTED, ERROR})

OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"  # encrypted Office files are OLE containers

TEXT_EXTENSIONS = {".txt", ".md", ".csv"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg"}
WORD_RE = re.compile(r"\b[^\W\d_]{3,}\b")  # words of 3+ letters, any script
SPREADSHEET_EXTENSIONS = {".xlsx", ".xls"}
SPREADSHEET_MAX_ROWS = 2000
OCR_MAX_IMAGE_WIDTH = 2500  # pixels; wider photos are read scaled down


@dataclass
class Extraction:
    status: str
    text: str = ""
    method: str = ""
    page_count: int | None = None
    truncated: bool = False
    error: str | None = None


def _clean(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \f\v]*\t[ \t\f\v]*", "\t", text)  # tabs separate table cells; keep them
    text = re.sub(r"[ \f\v]+", " ", text)
    text = re.sub(r"[ \t]*\n[ \t]*", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _finish(text: str, method: str, settings: ExtractionSettings, *,
            page_count: int | None = None, truncated: bool = False) -> Extraction:
    text = _clean(text)
    if len(text) > settings.max_chars:
        text, truncated = text[: settings.max_chars], True
    return Extraction(
        status=OK if text else EMPTY, text=text, method=method,
        page_count=page_count, truncated=truncated,
    )


def _tessdata_candidates(configured: str | None) -> list[Path]:
    candidates = [Path(configured)] if configured else []
    if os.environ.get("TESSDATA_PREFIX"):
        candidates.append(Path(os.environ["TESSDATA_PREFIX"]))
    exe = shutil.which("tesseract")
    if exe:
        candidates += [Path(exe).parent / "tessdata", Path(exe).parent.parent / "share" / "tessdata"]
    for base in filter(None, [os.environ.get("ProgramFiles"), os.environ.get("ProgramFiles(x86)"),
                              os.environ.get("LOCALAPPDATA") and str(Path(os.environ["LOCALAPPDATA"]) / "Programs")]):
        candidates.append(Path(base) / "Tesseract-OCR" / "tessdata")
    candidates += [Path("/opt/homebrew/share/tessdata"), Path("/usr/local/share/tessdata"),
                   Path("/usr/share/tessdata")]
    candidates += sorted(Path("/usr/share/tesseract-ocr").glob("*/tessdata"), reverse=True)
    return candidates


@lru_cache(maxsize=4)
def tessdata_path(configured: str | None = None) -> str | None:
    """Location of Tesseract language data, or None if Tesseract is not installed."""
    for folder in _tessdata_candidates(configured):
        try:
            if folder.is_dir() and any(folder.glob("*.traineddata")):
                return str(folder)
        except OSError:
            continue
    return None


def ocr_available(settings: ExtractionSettings) -> bool:
    return settings.ocr_enabled and tessdata_path(settings.tessdata_dir) is not None


@contextmanager
def _quiet_native_stderr():
    """Tesseract prints diagnostics ("Line cannot be recognized!!") straight to the console."""
    try:
        sys.stderr.flush()
        saved = os.dup(2)
    except (OSError, ValueError, AttributeError):  # no real stderr (e.g. under pythonw)
        yield
        return
    try:
        with open(os.devnull, "w") as devnull:
            os.dup2(devnull.fileno(), 2)
            yield
    finally:
        os.dup2(saved, 2)
        os.close(saved)


def _ocr_pages(doc, settings: ExtractionSettings, dpi: int = 300) -> str:
    tessdata = tessdata_path(settings.tessdata_dir)
    parts = []
    with _quiet_native_stderr():
        for page in doc.pages(0, min(doc.page_count, settings.ocr_max_pages)):
            textpage = page.get_textpage_ocr(language=settings.ocr_language, dpi=dpi, full=True, tessdata=tessdata)
            parts.append(page.get_text(textpage=textpage))
    return "\n".join(parts)


def _pdf(path: Path, settings: ExtractionSettings) -> Extraction:
    import pymupdf

    with pymupdf.open(path) as doc:
        if doc.needs_pass:
            return Extraction(ENCRYPTED, method="pymupdf", page_count=doc.page_count)
        pages = min(doc.page_count, settings.max_pages)
        text = "\n".join(page.get_text() for page in doc.pages(0, pages))
        truncated = doc.page_count > pages

        if pages and len(text.strip()) / pages < settings.min_chars_per_page:
            if not ocr_available(settings):
                result = _finish(text, "pymupdf", settings, page_count=doc.page_count, truncated=truncated)
                result.status = NEEDS_OCR
                return result
            text = _ocr_pages(doc, settings)
            truncated = doc.page_count > settings.ocr_max_pages
            return _finish(text, "pymupdf+ocr", settings, page_count=doc.page_count, truncated=truncated)

        return _finish(text, "pymupdf", settings, page_count=doc.page_count, truncated=truncated)


def _image(path: Path, settings: ExtractionSettings) -> Extraction:
    if not ocr_available(settings):
        return Extraction(NEEDS_OCR, method="ocr")
    import pymupdf

    with pymupdf.open(path) as doc:
        # Read the image at its own pixel resolution (a fixed 300 dpi would upscale
        # phone photos several times and make OCR very slow), capped for huge photos.
        width_px = pymupdf.Pixmap(str(path)).width
        dpi = max(36, round(72 * min(width_px, OCR_MAX_IMAGE_WIDTH) / doc[0].rect.width))
        text = _ocr_pages(doc, settings, dpi=dpi)
    result = _finish(text, "ocr", settings)
    if result.status == OK and len(WORD_RE.findall(result.text)) < settings.min_ocr_words:
        # OCR finds specks of "text" in ordinary photos; without real words it's a photo.
        return Extraction(EMPTY, method="ocr", error="No readable text; looks like a photo")
    return result


def _docx(path: Path, settings: ExtractionSettings) -> Extraction:
    import docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return _finish("\n".join(parts), "python-docx", settings)


def _pptx(path: Path, settings: ExtractionSettings) -> Extraction:
    from pptx import Presentation

    presentation = Presentation(str(path))
    parts = []
    for number, slide in enumerate(presentation.slides, start=1):
        parts.append(f"[Slide {number}]")
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
    return _finish("\n".join(parts), "python-pptx", settings, page_count=len(presentation.slides))


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _spreadsheet(path: Path, settings: ExtractionSettings) -> Extraction:
    from python_calamine import CalamineWorkbook

    workbook = CalamineWorkbook.from_path(str(path))
    parts, truncated = [], False
    try:
        for name in workbook.sheet_names:
            parts.append(f"[Sheet: {name}]")
            rows = workbook.get_sheet_by_name(name).to_python(skip_empty_area=True)
            if len(rows) > SPREADSHEET_MAX_ROWS:
                rows, truncated = rows[:SPREADSHEET_MAX_ROWS], True
            parts.extend("\t".join(_cell(v) for v in row) for row in rows)
    finally:
        workbook.close()
    return _finish("\n".join(parts), "calamine", settings,
                   page_count=len(workbook.sheet_names), truncated=truncated)


def _plain_text(path: Path, settings: ExtractionSettings) -> Extraction:
    # Up to 4 bytes per character covers max_chars in any of the encodings below.
    limit = settings.max_chars * 4
    with open(path, "rb") as fh:
        raw = fh.read(limit + 1)
    truncated = len(raw) > limit
    raw = raw[:limit]

    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        text = raw.decode("utf-16", errors="replace")
    else:
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            if truncated and exc.start >= len(raw) - 3:  # read limit split a character
                text = raw.decode("utf-8-sig", errors="ignore")
            else:
                text = raw.decode("cp1252", errors="replace")
    return _finish(text, "text", settings, truncated=truncated)


def _rtf(path: Path, settings: ExtractionSettings) -> Extraction:
    from striprtf.striprtf import rtf_to_text

    raw = path.read_bytes().decode("latin-1")
    return _finish(rtf_to_text(raw, errors="ignore"), "striprtf", settings)


EXTRACTORS = {
    ".pdf": _pdf,
    ".docx": _docx,
    ".pptx": _pptx,
    ".rtf": _rtf,
    **{ext: _spreadsheet for ext in SPREADSHEET_EXTENSIONS},
    **{ext: _plain_text for ext in TEXT_EXTENSIONS},
    **{ext: _image for ext in IMAGE_EXTENSIONS},
}


def _is_ole_container(path: Path) -> bool:
    with open(path, "rb") as fh:
        return fh.read(len(OLE_MAGIC)) == OLE_MAGIC


def extract(path: str | Path, settings: ExtractionSettings) -> Extraction:
    """Extract text from ``path``. Never raises for problems with the file itself."""
    path = Path(path)
    ext = path.suffix.lower()
    extractor = EXTRACTORS.get(ext)
    if extractor is None:
        return Extraction(UNSUPPORTED, error=f"No extractor for {ext or 'files without extension'}")
    try:
        if path.stat().st_size > settings.max_file_mb * 1024 * 1024:
            return Extraction(TOO_LARGE, error=f"Larger than {settings.max_file_mb} MB")
        if ext in {".docx", ".pptx", ".xlsx"} and _is_ole_container(path):
            return Extraction(ENCRYPTED, error="Password-protected Office file")
        return extractor(path, settings)
    except Exception as exc:  # corrupt files raise all sorts of library-specific errors
        log.warning("Extraction failed for %s: %s", path, type(exc).__name__)
        return Extraction(ERROR, error=f"{type(exc).__name__}: {exc}"[:500])
