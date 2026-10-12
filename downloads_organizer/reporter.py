"""CSV and Excel reports."""

from __future__ import annotations

import csv
import json
import sqlite3
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter

from downloads_organizer import extractors
from downloads_organizer.scanner import ScanResult

INVENTORY_COLUMNS = [
    "original_path",
    "filename",
    "extension",
    "size_bytes",
    "size",
    "created_at",
    "modified_at",
    "supported",
]
ERROR_COLUMNS = ["path", "error"]


def human_size(size: int) -> str:
    value = float(size)
    for unit in ("B", "KB", "MB", "GB"):
        if value < 1024 or unit == "GB":
            return f"{value:.0f} {unit}" if unit == "B" else f"{value:.1f} {unit}"
        value /= 1024
    return f"{value:.1f} GB"


def _inventory_rows(result: ScanResult) -> list[list]:
    return [
        [r.original_path, r.filename, r.extension, r.size_bytes, human_size(r.size_bytes),
         r.created_at or "", r.modified_at, "yes" if r.supported else "no"]
        for r in result.records
    ]


def summarize(result: ScanResult) -> dict:
    total_size = sum(r.size_bytes for r in result.records)
    return {
        "root": result.root,
        "started_at": result.started_at,
        "finished_at": result.finished_at,
        "files": len(result.records),
        "supported_files": sum(r.supported for r in result.records),
        "total_size": human_size(total_size),
        "errors": len(result.errors),
        "skipped": dict(result.skipped),
        "by_extension": Counter(r.extension or "(none)" for r in result.records).most_common(),
    }


def _write_csv(path: Path, header: list[str], rows: list[list]) -> None:
    # utf-8-sig so Excel opens non-ASCII filenames correctly.
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.writer(fh)
        writer.writerow(header)
        writer.writerows(rows)


def _add_sheet(wb: Workbook, title: str, header: list[str], rows: list[list]) -> None:
    ws = wb.create_sheet(title)
    ws.append(header)
    for row in rows:
        # Extracted text can contain control characters that Excel rejects.
        ws.append([ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v for v in row])
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    for idx, name in enumerate(header, start=1):
        width = max([len(str(name))] + [len(str(r[idx - 1])) for r in rows[:500] if len(r) >= idx])
        ws.column_dimensions[get_column_letter(idx)].width = min(max(width + 2, 10), 80)


def write_inventory(
    result: ScanResult,
    reports_dir: str | Path,
    formats: tuple[str, ...] = ("csv", "xlsx"),
) -> list[Path]:
    """Write the inventory report(s) and return the created file paths."""
    reports_dir = Path(reports_dir)
    reports_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    rows = _inventory_rows(result)
    error_rows = [[e.path, e.error] for e in result.errors]
    written: list[Path] = []

    if "csv" in formats:
        path = reports_dir / f"inventory_{stamp}.csv"
        _write_csv(path, INVENTORY_COLUMNS, rows)
        written.append(path)
        if error_rows:
            path = reports_dir / f"inventory_{stamp}_errors.csv"
            _write_csv(path, ERROR_COLUMNS, error_rows)
            written.append(path)

    if "xlsx" in formats:
        summary = summarize(result)
        wb = Workbook()
        wb.remove(wb.active)
        summary_rows = [
            ["Scanned folder", summary["root"]],
            ["Started", summary["started_at"]],
            ["Finished", summary["finished_at"]],
            ["Files inventoried", summary["files"]],
            ["Supported files", summary["supported_files"]],
            ["Total size", summary["total_size"]],
            ["Errors", summary["errors"]],
            *[[f"Skipped: {reason}", count] for reason, count in sorted(summary["skipped"].items())],
            [],
            ["Extension", "Files"],
            *[[ext, count] for ext, count in summary["by_extension"]],
        ]
        _add_sheet(wb, "Summary", ["Item", "Value"], summary_rows)
        _add_sheet(wb, "Inventory", INVENTORY_COLUMNS, rows)
        _add_sheet(wb, "Errors", ERROR_COLUMNS, error_rows)
        path = reports_dir / f"inventory_{stamp}.xlsx"
        wb.save(path)
        written.append(path)

    return written


REVIEW_REASONS = {
    extractors.NEEDS_OCR: "Scanned PDF or image with no text layer (needs OCR)",
    extractors.ENCRYPTED: "Password protected",
    extractors.TOO_LARGE: "Larger than the extraction size limit",
    extractors.UNSUPPORTED: "File type cannot be read (e.g. legacy .doc)",
    extractors.ERROR: "Could not be read; possibly corrupt",
    extractors.EMPTY: "Opened, but no text found",
}

DOCUMENT_COLUMNS = [
    "id", "original_path", "filename", "extension", "size_bytes", "size", "modified_at",
    "supported", "extraction_status", "method", "pages", "characters", "truncated",
    "duplicate_group", "classification_status", "document_type", "category", "recommended_action",
    "confidence", "importance", "sensitivity", "summary", "reasons", "classified_by",
    "sha256", "error", "text_preview",
]
PREVIEW_CHARS = 150

_DOCUMENTS_QUERY = """
SELECT d.*, t.method, t.page_count, t.char_count, t.truncated,
       substr(t.text, 1, ?) AS preview, g.group_id
FROM documents d
LEFT JOIN extracted_text t ON t.sha256 = d.sha256 AND d.supported = 1
LEFT JOIN duplicate_groups g ON g.document_id = d.id AND g.duplicate_type = 'exact'
WHERE d.present = 1
ORDER BY d.original_path
"""


def write_document_report(conn: sqlite3.Connection, reports_dir: str | Path) -> list[Path]:
    """Phase 2 report: extraction results, exact duplicates and the manual-review queue."""
    reports_dir = Path(reports_dir)
    reports_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    docs = conn.execute(_DOCUMENTS_QUERY, (PREVIEW_CHARS,)).fetchall()

    rows = []
    for d in docs:
        payload = json.loads(d["classification_json"]) if d["classification_json"] else {}
        c = payload.get("classification", {})
        rows.append([
            d["id"], d["original_path"], d["filename"], d["extension"], d["size_bytes"],
            human_size(d["size_bytes"]), d["modified_at"], "yes" if d["supported"] else "no",
            d["extraction_status"], d["method"] or "", d["page_count"], d["char_count"],
            "yes" if d["truncated"] else "", d["group_id"],
            d["classification_status"] if d["supported"] else "", c.get("document_type", ""),
            c.get("category", ""), c.get("recommended_action", ""), c.get("confidence"),
            c.get("importance", ""), c.get("sensitivity", ""), c.get("summary", ""),
            "; ".join(c.get("reasons", [])), payload.get("model_used", ""),
            d["sha256"] or "", d["processing_error"] or "", " ".join((d["preview"] or "").split()),
        ])
    classified = [d for d in docs if d["supported"] and d["classification_status"] == "done"]

    groups: dict[int, list[sqlite3.Row]] = {}
    for d in docs:
        if d["group_id"] is not None:
            groups.setdefault(d["group_id"], []).append(d)
    duplicate_rows = []
    reclaimable_total = 0
    for group_id, members in sorted(
        groups.items(), key=lambda kv: kv[1][0]["size_bytes"] * (len(kv[1]) - 1), reverse=True
    ):
        reclaimable = members[0]["size_bytes"] * (len(members) - 1)
        reclaimable_total += reclaimable
        for i, d in enumerate(sorted(members, key=lambda m: m["modified_at"])):
            duplicate_rows.append([
                group_id, len(members) if i == 0 else "", human_size(reclaimable) if i == 0 else "",
                d["filename"], d["original_path"], human_size(d["size_bytes"]), d["modified_at"],
            ])

    review_rows = [
        [REVIEW_REASONS[d["extraction_status"]], d["filename"], d["original_path"],
         human_size(d["size_bytes"]), d["extraction_status"], d["processing_error"] or ""]
        for d in docs
        if d["supported"] and d["extraction_status"] in extractors.REVIEW_STATUSES
    ]

    statuses = Counter(d["extraction_status"] for d in docs)
    summary_rows = [
        ["Files", len(docs)],
        ["Readable document types", sum(d["supported"] for d in docs)],
        ["Exact duplicate groups", len(groups)],
        ["Files in duplicate groups", sum(len(m) for m in groups.values())],
        ["Space used by extra copies", human_size(reclaimable_total)],
        ["Needs manual review", len(review_rows)],
        ["Classified by Nemotron", len(classified)],
        [],
        ["Recommendation", "Files"],
        *[[k, n] for k, n in Counter(d["recommended_action"] for d in classified).most_common()],
        [],
        ["Category", "Files"],
        *[[k, n] for k, n in Counter(d["category"] for d in classified).most_common()],
        [],
        ["Extraction status", "Files"],
        *[[status, count] for status, count in statuses.most_common()],
    ]

    csv_path = reports_dir / f"documents_{stamp}.csv"
    _write_csv(csv_path, DOCUMENT_COLUMNS, rows)

    wb = Workbook()
    wb.remove(wb.active)
    _add_sheet(wb, "Summary", ["Item", "Value"], summary_rows)
    _add_sheet(wb, "Duplicates", ["group", "copies", "extra space", "filename", "path", "size", "modified_at"],
               duplicate_rows)
    _add_sheet(wb, "Needs Review", ["reason", "filename", "path", "size", "status", "error"], review_rows)
    _add_sheet(wb, "Documents", DOCUMENT_COLUMNS, rows)
    xlsx_path = reports_dir / f"documents_{stamp}.xlsx"
    wb.save(xlsx_path)
    return [csv_path, xlsx_path]


_DECISIONS_QUERY = """
SELECT d.id, d.original_path, d.filename, d.size_bytes, d.modified_at, d.confidence, d.document_type,
       x.action, x.destination, x.ai_action, x.rules, x.reasons, p.original_path AS duplicate_of_path
FROM decisions x
JOIN documents d ON d.id = x.document_id
LEFT JOIN documents p ON p.id = x.duplicate_of
WHERE d.present = 1
ORDER BY x.action, x.destination, d.original_path
"""


def write_review_report(conn: sqlite3.Connection, reports_dir: str | Path) -> Path:
    """Phase 4 report: what the organizer proposes for each file, and why."""
    reports_dir = Path(reports_dir)
    reports_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    rows = conn.execute(_DECISIONS_QUERY).fetchall()

    def reasons(r) -> str:
        return " | ".join(json.loads(r["reasons"]))

    by_action = {}
    for r in rows:
        by_action.setdefault(r["action"], []).append(r)

    moves = [r for r in rows if r["action"] in ("KEEP", "ARCHIVE")]
    overridden = [r for r in rows if r["ai_action"] and r["ai_action"] != r["action"]]
    summary = [
        ["Files considered", len(rows)],
        *[[f"Decision: {action}", len(items)] for action, items in sorted(by_action.items())],
        ["Rules changed the AI's suggestion", len(overridden)],
        [],
        ["Destination (KEEP / ARCHIVE)", "Files"],
        *[[dest, n] for dest, n in Counter(r["destination"] for r in moves).most_common()],
        [],
        ["Why files are left in place", "Files"],
        *[[reason, n] for reason, n in Counter(
            json.loads(r["reasons"])[0] for r in by_action.get("LEAVE", [])).most_common()],
    ]

    wb = Workbook()
    wb.remove(wb.active)
    _add_sheet(wb, "Summary", ["Item", "Value"], summary)
    _add_sheet(
        wb, "Proposed Moves", ["action", "filename", "destination", "type", "confidence", "reason", "path"],
        [[r["action"], r["filename"], r["destination"], r["document_type"], r["confidence"], reasons(r),
          r["original_path"]] for r in moves],
    )
    _add_sheet(
        wb, "Review Queue", ["filename", "why", "AI suggested", "confidence", "size", "path"],
        [[r["filename"], reasons(r), r["ai_action"] or "", r["confidence"], human_size(r["size_bytes"]),
          r["original_path"]] for r in by_action.get("REVIEW", [])],
    )
    _add_sheet(
        wb, "Cleanup Candidates", ["filename", "why", "size", "kept copy", "path"],
        [[r["filename"], reasons(r), human_size(r["size_bytes"]), r["duplicate_of_path"] or "", r["original_path"]]
         for r in by_action.get("CLEANUP_CANDIDATE", [])],
    )
    _add_sheet(
        wb, "Rule Overrides", ["filename", "AI suggested", "final decision", "rule", "why", "path"],
        [[r["filename"], r["ai_action"], r["action"], ", ".join(json.loads(r["rules"])), reasons(r),
          r["original_path"]] for r in overridden],
    )
    _add_sheet(
        wb, "Left In Place", ["filename", "why", "path"],
        [[r["filename"], reasons(r), r["original_path"]] for r in by_action.get("LEAVE", [])],
    )
    path = reports_dir / f"review_{stamp}.xlsx"
    wb.save(path)
    return path


def write_move_plan(moves, skipped, reports_dir: str | Path) -> Path:
    """CSV preview of a move plan, so it can be checked before anything is moved."""
    reports_dir = Path(reports_dir)
    reports_dir.mkdir(parents=True, exist_ok=True)
    path = reports_dir / f"move_plan_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
    rows = [[m.document_id, m.action, str(m.source), str(m.destination), m.reason] for m in moves]
    rows += [[s.document_id, "SKIPPED", s.path, "", s.reason] for s in skipped]
    _write_csv(path, ["document_id", "action", "from", "to", "reason"], rows)
    return path
